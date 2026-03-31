#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
刑事法援会见前打印文档生成器

用法:
    python create_criminal_legal_aid_meeting.py [受援人] [身份证号] [住所/羁押地] [涉嫌罪名] [案件阶段] [代理权限]

示例:
    python create_criminal_legal_aid_meeting.py 张三 430124198203045123 宁乡市看守所 盗窃罪 侦查 一般代理
"""

import sys
from docx import Document
import shutil
import os


def create_criminal_legal_aid_meeting(weitr, wtrsfz, wtrzz, anyou, jied, quanx="刑事辩护", hjdd=""):
    """
    生成刑事法援会见前打印文档

    Args:
        weitr: 受援人姓名
        wtrsfz: 受援人身份证号
        wtrzz: 受援人住所/羁押地
        anyou: 涉嫌罪名
        jied: 案件阶段（侦查/审查起诉/审判）
        quanx: 代理权限（默认刑事辩护）
        hjdd: 会见地点
    """
    script_dir = os.path.dirname(os.path.abspath(__file__))
    project_dir = os.path.dirname(script_dir)
    template_path = os.path.join(project_dir, 'legal_templates', '8 刑事法援会见前打印.docx')
    output_dir = os.getcwd()

    safe_name = weitr.replace('/', '_').replace('\\', '_')
    output_path = os.path.join(output_dir, f'刑事法援会见前打印_{safe_name}_{anyou}_{jied}.docx')

    if os.path.exists(output_path):
        try:
            os.remove(output_path)
        except:
            output_path = os.path.join(output_dir, f'刑事法援会见前打印_{safe_name}_{anyou}_{jied}_new.docx')

    shutil.copy(template_path, output_path)
    doc = Document(output_path)

    placeholders = {
        'weitr': weitr,
        'wtrsfz': wtrsfz,
        'wtrzz': wtrzz,
        'anyou': anyou,
        'jied': jied,
        'wtxm': '一',
        'quanx': quanx,
        'hjdd': hjdd
    }

    for para in doc.paragraphs:
        full_text = para.text

        for placeholder, value in placeholders.items():
            if placeholder in full_text:
                for run in para.runs:
                    if placeholder in run.text:
                        run.text = run.text.replace(placeholder, value)

    doc.save(output_path)
    print(f'✅ 刑事法援会见前打印文档已生成: {output_path}')
    return output_path


if __name__ == '__main__':
    if len(sys.argv) >= 7:
        _, weitr, wtrsfz, wtrzz, anyou, jied, quanx = sys.argv[:7]
        hjdd = sys.argv[7] if len(sys.argv) > 7 else ""
        create_criminal_legal_aid_meeting(weitr, wtrsfz, wtrzz, anyou, jied, quanx, hjdd)
    else:
        print(__doc__)
        print("\n📝 使用默认数据生成...")
        print("-" * 50)

        weitr = "李四"
        wtrsfz = "430124198203045123"
        wtrzz = "长沙市第一看守所"
        anyou = "盗窃罪"
        jied = "侦查"
        quanx = "刑事辩护"
        hjdd = "长沙市第一看守所"

        print(f"\n使用数据:\n受援人: {weitr}\n身份证: {wtrsfz}\n住所/羁押地: {wtrzz}\n涉嫌罪名: {anyou}\n案件阶段: {jied}\n代理权限: {quanx}\n会见地点: {hjdd}")
        print("-" * 50)

        create_criminal_legal_aid_meeting(weitr, wtrsfz, wtrzz, anyou, jied, quanx, hjdd)
